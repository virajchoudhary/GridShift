"""Offline integrity checks for GridShift's frozen final artifacts."""

from __future__ import annotations

import gzip
import hashlib
import importlib
import importlib.util
import json
import math
import os
import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

SLOT_NS = 30 * 60 * 1_000_000_000
BASELINES = {"calendar_mean", "persistence", "daily_naive", "weekly_naive"}
BOUNDS = ("lower80", "upper80", "lower95", "upper95")
REQS = (
    ("R01", "Official data", ("identity", "data")),
    ("R02", "EDA", ("diagnostics",)), ("R03", "Missing values", ("data", "diagnostics")),
    ("R04", "ADF", ("diagnostics",)), ("R05", "Transformations", ("diagnostics", "forecasts")),
    ("R06", "STL", ("diagnostics",)), ("R07", "ACF/PACF", ("diagnostics",)),
    ("R08", "AR", ("diagnostics", "notebook")), ("R09", "MA", ("diagnostics", "notebook")),
    ("R10", "ARMA", ("diagnostics", "notebook")), ("R11", "ARIMA", ("diagnostics", "forecasts")),
    ("R12", "SARIMA", ("diagnostics", "forecasts")), ("R13", "Lag selection", ("diagnostics",)),
    ("R14", "AIC/BIC", ("diagnostics",)), ("R15", "Forecast comparison", ("identity", "forecasts")),
    ("R16", "Scheduling", ("primary", "queue")), ("R17", "Dashboard", ("dashboard", "tests")),
    ("R18", "Notebook", ("notebook",)), ("R19", "Report", ("exports",)),
    ("R20", "Presentation", ("exports",)), ("R21", "Local deployment", ("identity", "deployment")),
    ("R22", "Tests", ("tests",)), ("R23", "Reproducibility", ("identity", "data", "forecasts")),
)
REVIEW_REQUIRED = {f"R{n:02d}" for n in range(2, 15)} | {"R17", "R19", "R20"}


class MissingEvidence(Exception):
    pass


def _json_bytes(value):
    return json.dumps(value, indent=2, sort_keys=True, default=str, allow_nan=False).encode()


def _hash(value):
    return hashlib.sha256(value).hexdigest()


def _file_hash(path):
    return _hash(Path(path).read_bytes())


def _json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _safe(root, relative):
    path = Path(relative)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError(f"Unsafe project-relative path: {relative}")
    resolved, root = (root / path).resolve(), root.resolve()
    if resolved != root and root not in resolved.parents:
        raise ValueError(f"Path escapes project root: {relative}")
    return resolved


def _finite(value):
    try:
        return math.isfinite(float(value))
    except (TypeError, ValueError):
        return False


def _close(actual, expected, atol=1e-5, rtol=1e-7):
    return math.isclose(float(actual), float(expected), abs_tol=atol, rel_tol=rtol)


def _key(value):
    return None if pd.isna(value) else str(value)


def _software_hash(root):
    path = root / "src/gridshift/artifacts.py"
    spec = importlib.util.spec_from_file_location("_gridshift_artifacts_verify", path)
    if spec is None or spec.loader is None:
        return None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.software_digest(root)


def _identity(root, config):
    paths = [root / p for p in (
        "data/manifest.json", "data/normalized/carbon.parquet",
        "artifacts/protocol_freeze.json", "artifacts/results.json",
    )]
    missing = [str(p.relative_to(root)) for p in paths if not p.is_file()]
    if missing:
        raise MissingEvidence("Missing identity inputs: " + ", ".join(missing))
    manifest, freeze, results = map(_json, (paths[0], paths[2], paths[3]))
    protocol, data_hash = freeze.get("protocol"), manifest.get("normalized_hash")
    if not isinstance(protocol, dict) or not isinstance(data_hash, str):
        raise ValueError("Invalid freeze or data manifest")
    if manifest.get("units") != "gCO2/kWh" or manifest.get("target") != "NESO estimated actual":
        raise ValueError("Data manifest target/units differ from the approved scope")
    protocol_hash = _hash(_json_bytes(protocol))
    study_id = _hash(_json_bytes({"data": data_hash, "config": config}))[:16]
    if _file_hash(paths[1]) != data_hash or freeze.get("data_hash") != data_hash:
        raise ValueError("Normalized data hash differs from manifest/freeze")
    if freeze.get("protocol_hash") != protocol_hash or protocol.get("config") != config or protocol.get("study_id") != study_id:
        raise ValueError("Freeze hash, configuration, or study_id mismatch")
    software = protocol.get("software_hash")
    if not isinstance(software, str) or results.get("software_hash") != software:
        raise ValueError("Results software hash differs from freeze")
    if software != _software_hash(root):
        raise ValueError("Current scientific software differs from freeze")
    if (results.get("study_id"), results.get("data_hash"), results.get("protocol_hash")) != (study_id, data_hash, protocol_hash):
        raise ValueError("Results data/protocol/study identity differs from freeze")
    frozen_files = protocol.get("validation_artifact_hashes")
    if not isinstance(frozen_files, dict) or not frozen_files:
        raise ValueError("Freeze has no validation artifact hashes")
    for relative, expected in frozen_files.items():
        path = _safe(root, relative)
        if not path.is_file() or _file_hash(path) != expected:
            raise ValueError(f"Frozen validation artifact changed: {relative}")
    return {
        "data_hash": data_hash, "protocol_hash": protocol_hash, "study_id": study_id,
        "software_hash": software, "run_dir": f"artifacts/runs/{study_id}",
        "deployment_model": protocol.get("deployment_model"),
    }


def _data(root, identity):
    if not (root / "data/manifest.json").is_file() or not (root / "data/normalized/carbon.parquet").is_file():
        raise MissingEvidence("Data manifest or normalized dataset is missing")
    manifest = _json(root / "data/manifest.json")
    chunks = manifest.get("chunks")
    if not isinstance(chunks, list) or not chunks:
        raise ValueError("Data manifest has no raw chunk hashes")
    for chunk in chunks:
        path = _safe(root, chunk["path"])
        if not path.is_file() or _hash(gzip.decompress(path.read_bytes())) != chunk.get("raw_hash"):
            raise ValueError(f"Raw source hash mismatch: {chunk['path']}")
    quality_path, reconstruction_path = root / "artifacts/data/quality.json", root / "artifacts/data/reconstruction.json"
    if not quality_path.is_file() or not reconstruction_path.is_file():
        raise MissingEvidence("Data quality/reconstruction evidence is missing")
    quality, reconstruction = map(_json, (quality_path, reconstruction_path))
    frame = pd.read_parquet(root / "data/normalized/carbon.parquet")
    actual = pd.to_numeric(frame["actual"], errors="raise").to_numpy(dtype=float)
    observed = np.isfinite(actual)
    if np.isinf(actual).any() or np.any(actual[observed] < 0):
        raise ValueError("Normalized actuals contain invalid values")
    counts = (len(frame), int(observed.sum()), int((~observed).sum()), len(chunks))
    if (
        quality.get("slots") != counts[0] or quality.get("observed") != counts[1]
        or quality.get("missing") != counts[2] or reconstruction.get("reconstructed_slots") != counts[0]
        or reconstruction.get("raw_chunks") != counts[3] or reconstruction.get("values_match") is not True
        or reconstruction.get("raw_hashes_match") is not True
    ):
        raise ValueError("Quality and reconstruction counts do not match normalized/raw data")
    return {"raw_chunks": counts[3], "normalized_slots": counts[0], "observed": counts[1], "data_hash": identity["data_hash"]}


def _forecast_file(path, root, origins, horizon, delay, require_provenance=False):
    sidecar = path.with_suffix(".json")
    if not path.is_file() or not sidecar.is_file():
        raise MissingEvidence(f"Forecast NPZ/sidecar missing: {path.name}")
    metadata = _json(sidecar)
    model_id = metadata.get("model_id", metadata.get("model_spec", {}).get("name"))
    if require_provenance:
        if not model_id:
            raise ValueError(f"Missing forecast model provenance: {path.name}")
        if model_id not in BASELINES and (
            "publication_delay_slots" not in metadata
            or not isinstance(metadata.get("refits"), list) or not metadata["refits"]
        ):
            raise ValueError(f"Missing forecast cutoff provenance: {path.name}")
    if metadata.get("forecast_sha256") != _file_hash(path):
        raise ValueError(f"Forecast NPZ checksum mismatch: {path.name}")
    with np.load(path, allow_pickle=False) as arrays:
        required = {"origins_ns", "mean", "std", "fallback", *BOUNDS}
        if not required.issubset(arrays.files):
            raise ValueError(f"Forecast NPZ arrays missing: {path.name}")
        saved_origins = np.asarray(arrays["origins_ns"])
        if saved_origins.dtype != np.dtype("int64") or not np.array_equal(saved_origins, origins):
            raise ValueError(f"Forecast origin grid mismatch: {path.name}")
        if arrays["fallback"].dtype != np.dtype("bool") or arrays["fallback"].shape != (len(origins),):
            raise ValueError(f"Invalid fallback flags: {path.name}")
        values = [np.asarray(arrays[key]) for key in ("mean", "std", *BOUNDS)]
        if any(v.shape != (len(origins), horizon) or v.dtype.kind != "f" or not np.isfinite(v).all() for v in values):
            raise ValueError(f"Non-finite or wrong-horizon forecasts: {path.name}")
        mean, std, low80, high80, low95, high95 = values
        if np.any(std < 0) or np.any(low95 > low80) or np.any(low80 > high80) or np.any(high80 > high95):
            raise ValueError(f"Invalid forecast intervals: {path.name}")
    if metadata.get("origin_count", len(origins)) != len(origins):
        raise ValueError(f"Forecast origin count mismatch: {path.name}")
    if "horizon_slots" in metadata and int(metadata["horizon_slots"]) != horizon:
        raise ValueError(f"Forecast horizon metadata mismatch: {path.name}")
    if "publication_delay_slots" in metadata and int(metadata["publication_delay_slots"]) != delay:
        raise ValueError(f"Forecast delay metadata mismatch: {path.name}")
    cursor = 0
    for refit in metadata.get("refits", []):
        origin = pd.Timestamp(refit["refit_origin"]).tz_convert("UTC")
        cutoff = pd.Timestamp(refit["training_cutoff"]).tz_convert("UTC")
        count = int(refit["origins"])
        if count < 1 or cursor + count > len(origins) or origin.value != int(origins[cursor]):
            raise ValueError(f"Forecast refit partition mismatch: {path.name}")
        if cutoff.value != origin.value - (1 + delay) * SLOT_NS:
            raise ValueError(f"Forecast training cutoff mismatch: {path.name}")
        cursor += count
    if metadata.get("refits") and cursor != len(origins):
        raise ValueError(f"Forecast refits do not cover all origins: {path.name}")
    return str(path.relative_to(root))


def forecast_preflight(root, config, stage="test"):
    """Check official forecast provenance before executing either backtest."""
    import tempfile
    from .artifacts import require_freeze
    from .baselines import rolling
    from .data import load
    from .study import forecast_config, run_dir

    root = Path(root).resolve()
    frozen = require_freeze(root, config)
    directory = run_dir(root, config, stage)
    mapping = _json(directory / "forecast_store.json")
    study = config["study"]
    start, end = pd.Timestamp(study[stage + "_start"]), pd.Timestamp(study[stage + "_end"])
    horizon, delay = study["horizon_slots"], study["publication_delay_slots"]
    origins = pd.date_range(start, end - pd.Timedelta(minutes=30 * horizon), freq="30min").asi8
    specifications = {s["name"]: s for s in frozen["protocol"]["family_finalists"]}
    required = set(specifications) | BASELINES
    if not required.issubset(mapping):
        raise ValueError("Forecast preflight lacks a frozen model")
    hashes = {}
    for name in sorted(required):
        path = _safe(root, mapping[name])
        _forecast_file(path, root, origins, horizon, delay, require_provenance=True)
        metadata = _json(path.with_suffix(".json"))
        if metadata.get("model_id", metadata.get("model_spec", {}).get("name")) != name:
            raise ValueError("Forecast preflight model identity mismatch")
        if name in specifications and metadata.get("model_spec") != specifications[name]:
            raise ValueError("Forecast preflight specification differs from frozen finalist")
        hashes[mapping[name]] = _file_hash(path)
    with tempfile.TemporaryDirectory(prefix="gridshift-baseline-proof-") as temporary:
        rebuilt = rolling(load(root).actual, start, end, temporary, forecast_config(config), experiment_root=root)
        for name, regenerated in rebuilt.items():
            with np.load(regenerated, allow_pickle=False) as expected, np.load(
                _safe(root, mapping[name]), allow_pickle=False,
            ) as saved:
                if set(saved.files) != set(expected.files):
                    raise ValueError(f"Baseline array schema mismatch: {name}")
                for key in expected.files:
                    if not np.array_equal(saved[key], expected[key], equal_nan=True):
                        raise ValueError(f"Baseline causal reconstruction mismatch: {name} {key}")
    return {
        "passed": True, "stage": stage, "models": len(required), "origins": len(origins),
        "protocol_hash": frozen["protocol_hash"], "data_hash": frozen["data_hash"],
        "software_hash": frozen["protocol"]["software_hash"], "forecast_hashes": hashes,
        "classical_cutoffs": "Complete checked refit partitions; exact eligible cutoff and explicit delay",
        "baseline_cutoffs": "All arrays reconstructed from frozen source and causal eligible-prefix algorithm",
    }


def _forecasts(root, config, identity):
    stage = root / identity["run_dir"] / "test"
    store_path = stage / "forecast_store.json"
    if not store_path.is_file():
        raise MissingEvidence("Test forecast store is missing")
    mapping = _json(store_path)
    protocol = _json(root / "artifacts/protocol_freeze.json")["protocol"]
    models = {s["name"] for s in protocol.get("family_finalists", [])} | BASELINES
    if not models.issubset(mapping):
        raise MissingEvidence("Test store lacks a frozen finalist or baseline")
    study = config["study"]
    start, end = pd.Timestamp(study["test_start"]), pd.Timestamp(study["test_end"])
    horizon, delay = int(study["horizon_slots"]), int(study["publication_delay_slots"])
    origins = pd.date_range(start, end - pd.Timedelta(minutes=30 * horizon), freq="30min").asi8.astype(np.int64)
    files = [_safe(root, mapping[name]) for name in sorted(models)]
    verified = [_forecast_file(path, root, origins, horizon, delay, require_provenance=True) for path in files]
    forecast_preflight(root, config)
    return {
        "models": len(verified), "origins": len(origins), "horizon_slots": horizon,
        "cutoff_rule": f"origin minus {1 + delay} half-hour slots", "files": verified,
    }


def _independent_cost(actuals, start, duration, power, pue):
    index, origin = pd.DatetimeIndex(actuals.index), pd.Timestamp(start)
    if index.tz is None or origin.tzinfo is None or origin.value % SLOT_NS:
        raise ValueError("Actuals/start must use timezone-aware half-hour timestamps")
    duration, power, pue = float(duration), float(power), float(pue)
    if not all(map(math.isfinite, (duration, power, pue))) or duration <= 0 or power <= 0 or pue < 1:
        raise ValueError("Invalid workload energy assumptions")
    count = math.ceil(duration * 2 - 1e-12)
    overlap = np.minimum(0.5, np.maximum(0, duration - np.arange(count) * 0.5))
    slots = pd.date_range(origin.tz_convert("UTC"), periods=count, freq="30min")
    values = pd.Series(pd.to_numeric(actuals, errors="raise").to_numpy(float), index=index.tz_convert("UTC")).reindex(slots).to_numpy()
    if not np.isfinite(values).all() or np.any(values < 0):
        raise ValueError("Actual intensity slice is missing or invalid")
    return float(np.dot(overlap * power * pue, values))


def _primary(root, config, identity):
    directory = root / identity["run_dir"] / "test/single_job"
    dp, pp = directory / "primary_decisions.parquet", directory / "primary_pairs.parquet"
    if not dp.is_file() or not pp.is_file():
        raise MissingEvidence("Primary decisions/pairs are missing")
    decisions, pairs = pd.read_parquet(dp), pd.read_parquet(pp)
    dcols = {"scenario_id", "job_id", "release", "deadline", "duration_hours", "power_kw", "pue", "policy", "model_id", "status", "start", "finish", "realized_grams"}
    pcols = {"scenario_id", "job_id", "policy", "model_id", "baseline_grams", "policy_grams", "savings_grams", "savings_pct"}
    keys = ["scenario_id", "job_id", "policy", "model_id"]
    if not dcols.issubset(decisions) or not pcols.issubset(pairs):
        raise ValueError("Primary table schema is incomplete")
    if decisions.duplicated(keys).any() or pairs.duplicated(keys).any():
        raise ValueError("Duplicate primary decision/pair keys")
    if not set(decisions.status).issubset({"scheduled", "infeasible", "unavailable"}):
        raise ValueError("Unknown primary status")
    scheduled = decisions.loc[decisions.status == "scheduled"]
    if not scheduled.realized_grams.map(_finite).all():
        raise ValueError("Scheduled primary work lacks realized emissions")
    for row in scheduled.itertuples():
        times = [pd.Timestamp(v) for v in (row.release, row.deadline, row.start, row.finish)]
        if any(t.tzinfo is None for t in times):
            raise ValueError(f"Naive primary timestamp: {row.job_id}")
        release, deadline, start, finish = times
        if start < release or finish > deadline or finish <= start or not _close((finish - start).total_seconds() / 3600, row.duration_hours):
            raise ValueError(f"Primary feasibility violation: {row.job_id}")
    lookup, baseline = {}, {}
    for row in decisions.itertuples():
        key = (str(row.scenario_id), str(row.job_id), str(row.policy), _key(row.model_id))
        lookup[key] = row
        if row.policy == "earliest":
            baseline[(key[0], key[1])] = row
    expected = set()
    for key, row in lookup.items():
        base = baseline.get(key[:2])
        if key[2] != "earliest" and base is not None and row.status == base.status == "scheduled":
            expected.add(key)
    compared = set()
    for row in pairs.itertuples():
        key = (str(row.scenario_id), str(row.job_id), str(row.policy), _key(row.model_id))
        policy, base = lookup.get(key), baseline.get(key[:2])
        if policy is None or base is None or not _finite(row.baseline_grams) or not _finite(row.policy_grams):
            raise ValueError("Primary pair has no finite matching decisions")
        saved = float(row.baseline_grams) - float(row.policy_grams)
        pct = 100 * saved / float(row.baseline_grams) if float(row.baseline_grams) > 0 else np.nan
        if not (_close(row.baseline_grams, base.realized_grams) and _close(row.policy_grams, policy.realized_grams) and _close(row.savings_grams, saved)):
            raise ValueError("Primary pair values differ from decisions")
        if math.isfinite(pct) and not _close(row.savings_pct, pct):
            raise ValueError("Primary paired savings percentage does not recompute")
        compared.add(key)
    if compared != expected:
        raise ValueError("Primary compared job/model sets differ from finite matched decisions")
    for scenario_id, scenario in decisions.groupby("scenario_id"):
        oracle = scenario.loc[scenario.policy == "oracle"].set_index("job_id")
        for row in scenario.loc[(scenario.policy != "oracle") & (scenario.status == "scheduled")].itertuples():
            if row.job_id not in oracle.index or oracle.loc[row.job_id].status != "scheduled":
                raise ValueError(f"Missing oracle decision: {scenario_id}/{row.job_id}")
            value = float(oracle.loc[row.job_id].realized_grams)
            if not math.isfinite(value) or value > float(row.realized_grams) + 1e-5:
                raise ValueError(f"Oracle exceeds/misses realized cost: {scenario_id}/{row.job_id}")
    default_id = f"duration_{float(config['scheduling']['default_duration_hours']):g}h_slack_{float(config['scheduling']['default_slack_hours']):g}h"
    default = decisions.loc[decisions.scenario_id == default_id]
    selected = default.loc[(default.policy == "forecast") & (default.model_id == identity["deployment_model"])]
    earliest = default.loc[default.policy == "earliest"]
    if not len(selected) or set(selected.job_id) != set(earliest.job_id) or not (selected.status == "scheduled").all():
        raise ValueError("Default deployment/earliest job sets differ or contain incomplete work")
    if _json(root / "artifacts/results.json").get("primary_jobs") != len(selected):
        raise ValueError("results.json primary job count differs from selected decisions")
    normalized = pd.read_parquet(root / "data/normalized/carbon.parquet")
    actuals = pd.Series(pd.to_numeric(normalized["actual"], errors="raise").to_numpy(float), index=normalized.index)
    sample = default.loc[(default.status == "scheduled") & default.realized_grams.map(_finite)].sort_values(["release", "policy", "model_id"], na_position="first")
    if sample.empty:
        raise MissingEvidence("No realized primary rows are available for an independent sample")
    positions = np.unique(np.linspace(0, len(sample) - 1, min(64, len(sample)), dtype=int))
    errors = [
        abs(_independent_cost(actuals, r.start, r.duration_hours, r.power_kw, r.pue) - float(r.realized_grams))
        for r in sample.iloc[positions].itertuples()
    ]
    if max(errors) > 1e-4:
        raise ValueError("Independent realized-grams sample differs from saved primary values")
    return {"decisions": len(decisions), "pairs": len(pairs), "jobs": len(selected), "sample": len(errors), "max_error_grams": max(errors)}


def _reservations(rows):
    intervals = []
    for row in rows.loc[rows.status == "scheduled"].itertuples():
        release, deadline, start, finish = [pd.Timestamp(x) for x in (row.release, row.deadline, row.start, row.finish)]
        if any(x.tzinfo is None for x in (release, deadline, start, finish)) or start < release or finish > deadline or finish <= start:
            raise ValueError(f"Infeasible queue reservation: {row.job_id}")
        if not _close((finish - start).total_seconds() / 3600, row.duration_hours):
            raise ValueError(f"Queue duration mismatch: {row.job_id}")
        intervals.append((start.tz_convert("UTC"), finish.tz_convert("UTC")))
    intervals.sort()
    if any(b[0] < a[1] for a, b in zip(intervals, intervals[1:])):
        raise ValueError("Overlapping queue reservations")
    return len(intervals)


def _queue(root, config, identity):
    directory = root / identity["run_dir"] / "test/queue"
    completion = directory / "completion.json"
    if not completion.is_file():
        raise MissingEvidence("Queue completion is missing")
    levels, seeds, slacks = config["queue"]["offered_loads"], config["queue"]["seeds"], config["queue"]["slack_hours"]
    stems = {f"load_{level}_seed_{seed}_slack_{slack}" for level in levels for seed in seeds for slack in slacks}
    if len(stems) != 18 or len(seeds) != 3 or (_json(completion).get("traces"), _json(completion).get("completed")) != (18, True):
        raise ValueError("Queue completion does not cover 18 traces and three seeds")
    summary_path, boot_path = directory / "queue_summary.csv", directory / "queue_bootstrap.csv"
    if not summary_path.is_file() or not boot_path.is_file():
        raise MissingEvidence("Queue summary/bootstrap is missing")
    summary, bootstrap = pd.read_csv(summary_path), pd.read_csv(boot_path)
    policies, reservation_count = {"earliest", "strongest_simple", "selected"}, 0
    for stem in stems:
        paths = [directory / f"{stem}_{suffix}.parquet" for suffix in ("arrivals", "decisions", "scores")]
        if not all(p.is_file() for p in paths):
            raise MissingEvidence(f"Incomplete queue trace: {stem}")
        arrivals, decisions, scores = map(pd.read_parquet, paths)
        if arrivals.job_id.duplicated().any() or decisions.duplicated(["job_id", "policy"]).any() or scores.duplicated(["job_id", "policy"]).any():
            raise ValueError(f"Duplicate queue work: {stem}")
        if set(decisions.policy) != policies or set(scores.policy) != policies:
            raise ValueError(f"Expected three queue policies: {stem}")
        offered = set(arrivals.job_id)
        for policy in policies:
            d, s = decisions.loc[decisions.policy == policy], scores.loc[scores.policy == policy]
            if set(d.job_id) != offered or set(s.job_id) != offered:
                raise ValueError(f"Queue omitted offered/rejected work: {stem}/{policy}")
            reservation_count += _reservations(d)
        if set(zip(decisions.job_id, decisions.policy)) != set(zip(scores.job_id, scores.policy)):
            raise ValueError(f"Queue scores omit decision rows: {stem}")
        match = re.fullmatch(r"load_([^_]+)_seed_(\d+)_slack_([^_]+)", stem)
        level, seed, slack = float(match[1]), int(match[2]), float(match[3])
        q = summary.loc[np.isclose(summary.offered_load.astype(float), level) & (summary.seed.astype(int) == seed) & np.isclose(summary.slack_hours.astype(float), slack)]
        if len(q) != 3 or set(q.policy) != policies:
            raise ValueError(f"Queue summary rows missing: {stem}")
        offered_rows = scores.loc[scores.scored_release.astype(bool)]
        accepted = {p: set(offered_rows.loc[(offered_rows.policy == p) & (offered_rows.status == "scheduled"), "job_id"]) for p in policies}
        equal = all(ids == accepted["earliest"] for ids in accepted.values())
        common = set.intersection(*accepted.values())
        matched = offered_rows.loc[offered_rows.job_id.isin(common) & offered_rows.full_window_observed.astype(bool)]
        if matched.realized_grams.map(_finite).sum() != len(matched):
            raise ValueError(f"Queue matched work has missing actuals: {stem}")
        base = float(matched.loc[matched.policy == "earliest", "realized_grams"].sum())
        label = "fixed accepted work" if equal else "descriptive completion intersection; service differs"
        for row in q.itertuples():
            work = offered_rows.loc[offered_rows.policy == row.policy]
            grams = float(matched.loc[matched.policy == row.policy, "realized_grams"].sum())
            percent = 100 * (base - grams) / base if base > 0 else np.nan
            if row.interpretation != label or bool(row.fixed_work_comparison) != equal:
                raise ValueError(f"Queue savings label misstates accepted work: {stem}/{row.policy}")
            if int(row.offered_jobs) != len(work) or int(row.completed_jobs) != int((work.status == "scheduled").sum()) or int(row.rejected_jobs) != int((work.status != "scheduled").sum()):
                raise ValueError(f"Queue summary omits offered/completed/rejected work: {stem}/{row.policy}")
            if int(row.common_completed_observed_jobs) != int((matched.policy == row.policy).sum()):
                raise ValueError(f"Queue common job count mismatch: {stem}/{row.policy}")
            if not (_close(row.matched_grams, grams) and _close(row.matched_baseline_grams, base) and _close(row.matched_savings_grams, base - grams)):
                raise ValueError(f"Queue matched totals do not recompute: {stem}/{row.policy}")
            if math.isfinite(percent) != _finite(row.matched_savings_pct) or (math.isfinite(percent) and not _close(row.matched_savings_pct, percent)):
                raise ValueError(f"Queue savings percentage does not recompute: {stem}/{row.policy}")
        b = bootstrap.loc[np.isclose(bootstrap.offered_load.astype(float), level) & (bootstrap.seed.astype(int) == seed) & np.isclose(bootstrap.slack_hours.astype(float), slack)]
        if b.empty or set(b.interpretation) != {label}:
            raise ValueError(f"Queue bootstrap label mismatch: {stem}")
    return {"traces": len(stems), "policies_per_trace": len(policies), "reservations": reservation_count}


def _tree_hash(path):
    if path.is_file():
        return _file_hash(path)
    if not path.is_dir():
        raise MissingEvidence(f"Missing export output: {path}")
    digest = hashlib.sha256()
    for child in sorted(p for p in path.rglob("*") if p.is_file()):
        digest.update(child.relative_to(path).as_posix().encode())
        digest.update(bytes.fromhex(_file_hash(child)))
    return digest.hexdigest()


def _exports(root, identity):
    manifest_path = root / "reports/export_manifest.json"
    if not manifest_path.is_file():
        raise MissingEvidence("Export manifest is missing")
    manifest = _json(manifest_path)
    if manifest.get("study_id") != identity["study_id"]:
        raise ValueError("Export belongs to another study")
    inputs = manifest.get("input_hashes", {})
    config = _json(root / "artifacts/protocol_freeze.json")["protocol"]["config"]
    if inputs.get("config_json") != _hash(_json_bytes(config)):
        raise ValueError("Export configuration hash differs from the freeze")
    generator = hashlib.sha256()
    for source in (root / "src/gridshift/exports.py", root / "src/gridshift/presentation.mjs"):
        if not source.is_file():
            raise MissingEvidence(f"Export generator source is missing: {source.name}")
        generator.update(source.name.encode("utf-8"))
        generator.update(source.read_bytes())
    if manifest.get("generator_hash") != generator.hexdigest():
        raise ValueError("Export generator hash differs from current source")
    outputs = [
        "reports/gridshift_report.md", "reports/gridshift_report.pdf", "reports/gridshift_report_preview",
        "notebooks/01_gridshift_study.ipynb", "notebooks/01_gridshift_study.html",
        "slides/gridshift_presentation_content.json", "slides/gridshift_presentation.mjs",
        "slides/gridshift_presentation.pptx", "slides/gridshift_presentation.pdf",
        *[f"slides/gridshift_presentation_preview/slide-{n:02d}.png" for n in range(1, 13)],
    ]
    for relative in outputs:
        if manifest.get("output_hashes", {}).get(relative) != _tree_hash(_safe(root, relative)):
            raise ValueError(f"Export hash mismatch: {relative}")
    from pptx import Presentation
    deck = Presentation(str(root / "slides/gridshift_presentation.pptx"))
    if len(deck.slides) != 12 or any(not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip() for slide in deck.slides):
        raise ValueError("PPTX must contain 12 slides with nonempty notes")
    if any(not (root / f).read_bytes().startswith(b"%PDF-") for f in ("reports/gridshift_report.pdf", "slides/gridshift_presentation.pdf")):
        raise ValueError("An academic PDF has an invalid signature")
    if manifest.get("presentation_slides") != 12 or manifest.get("presentation_slide_pngs") != 12:
        raise ValueError("Export receipt does not record 12 slides and previews")
    pages = manifest.get("report_total_pages")
    if not isinstance(pages, int) or pages < 1 or manifest.get("report_rendered_pages") != pages:
        raise ValueError("Report page/render counts do not agree")
    return {"outputs": len(outputs), "slides": 12, "notes": 12, "report_pages": pages, "presentation_pdf_pages": 12}


def _notebook(root):
    notebook, html = root / "notebooks/01_gridshift_study.ipynb", root / "notebooks/01_gridshift_study.html"
    if not notebook.is_file() or not html.is_file():
        raise MissingEvidence("Notebook or its HTML export is missing")
    doc = _json(notebook)
    code = [cell for cell in doc.get("cells", []) if cell.get("cell_type") == "code"]
    if doc.get("nbformat") != 4 or not code or any(c.get("execution_count") is None for c in code):
        raise MissingEvidence("Notebook code cells are not all executed")
    errors = [o for c in code for o in c.get("outputs", []) if o.get("output_type") == "error"]
    if errors:
        raise ValueError(f"Notebook contains {len(errors)} errors")
    export_manifest = root / "reports/export_manifest.json"
    if not export_manifest.is_file():
        raise MissingEvidence("Notebook execution receipt is missing")
    status = _json(export_manifest).get("notebook_execution")
    if status != "EXECUTED_OFFLINE_WITHOUT_ERRORS" or b"<html" not in html.read_bytes().lower():
        raise MissingEvidence("Notebook execution receipt or HTML export is incomplete")
    return {"code_cells": len(code), "errors": 0, "html": str(html.relative_to(root))}


def _deployment(root, identity):
    directory, path = root / "artifacts/deployment", root / "artifacts/deployment/deployment.json"
    if not path.is_file():
        raise MissingEvidence("Deployment manifest is missing")
    manifest = _json(path)
    if manifest.get("schema_version") != 1 or manifest.get("protocol_hash") != identity["protocol_hash"] or manifest.get("data_hash") != identity["data_hash"]:
        raise ValueError("Deployment manifest is incompatible with the freeze")
    rel = Path(manifest.get("file", ""))
    if rel.is_absolute() or ".." in rel.parts:
        raise ValueError("Deployment path escapes its directory")
    bundle = (directory / rel).resolve()
    if directory.resolve() not in bundle.parents:
        raise ValueError("Deployment payload escapes its trusted directory")
    if not bundle.is_file():
        raise MissingEvidence("Deployment payload is missing")
    if _file_hash(bundle) != manifest.get("file_hash"):
        raise ValueError("Deployment payload hash/path mismatch")
    if manifest.get("library_versions") != {"numpy": np.__version__, "pandas": pd.__version__}:
        raise ValueError("Deployment runtime versions differ")
    fixture_path = directory / "prediction_fixture.json"
    if not fixture_path.is_file():
        raise MissingEvidence("Deployment prediction fixture is missing")
    if _file_hash(fixture_path) != manifest.get("fixture_hash"):
        raise ValueError("Deployment prediction fixture hash mismatch")
    if manifest.get("kind") == "classical":
        model = _json(bundle)
        arrays = bundle.with_name(model.get("arrays_file", ""))
        if arrays.parent != bundle.parent or not arrays.is_file() or _file_hash(arrays) != model.get("arrays_sha256"):
            raise ValueError("Deployment model arrays fail checksum")
    elif manifest.get("kind") != "baseline":
        raise ValueError("Unsupported deployment bundle kind")
    src = str(root / "src")
    sys.path.insert(0, src)
    try:
        live = importlib.import_module("gridshift.live")
        fixture = live.verify_deployment_fixture(root, _json(root / "artifacts/protocol_freeze.json")["protocol"]["config"])
    finally:
        if sys.path and sys.path[0] == src:
            sys.path.pop(0)
    if fixture.get("verified") is not True or fixture.get("realized_result") != "not scored":
        raise ValueError("Deployment prediction fixture failed or was scored")
    return {"model": manifest["model_id"], "kind": manifest["kind"], "fixture": fixture}


def _tests(root, identity):
    junit, receipt_path = root / "artifacts/tests/core.xml", root / "artifacts/tests/core_receipt.json"
    if not junit.is_file():
        raise MissingEvidence("Current JUnit report is missing")
    doc = ET.parse(junit).getroot()
    suites = [doc] if doc.tag == "testsuite" else doc.findall(".//testsuite")
    if not suites:
        raise ValueError("JUnit report contains no suites")
    tests = sum(int(s.get("tests", 0)) for s in suites)
    failures = sum(int(s.get("failures", 0)) for s in suites)
    errors = sum(int(s.get("errors", 0)) for s in suites)
    if failures or errors:
        raise ValueError(f"Current suite has {failures} failures and {errors} errors in {tests} tests")
    if not receipt_path.is_file():
        raise MissingEvidence("Passing JUnit receipt is missing")
    receipt = _json(receipt_path)
    source = _software_hash(root)
    if (
        receipt.get("passed") is not True or receipt.get("tests") != tests
        or receipt.get("junit_sha256") != _file_hash(junit)
        or receipt.get("software_hash") != source
        or receipt.get("software_hash") != identity.get("software_hash")
    ):
        raise ValueError("JUnit receipt is stale or does not bind current source")
    if any(p.is_file() and p.stat().st_mtime > receipt_path.stat().st_mtime + 2 for p in [*root.glob("tests/test_*.py"), root / "scripts/test_project.py"]):
        raise ValueError("Tests changed after their passing receipt")
    return {"tests": tests, "failures": 0, "errors": 0, "software_hash": source}


DIAGNOSTICS = (
    "diagnostics.json", "distribution.csv", "daily_slot_heatmap.csv", "monthly_distribution.csv",
    "rolling_summary.csv", "weekday_local_profile.csv", "year_season.csv", "periodogram.csv",
    "stl48_components.csv", "acf_pacf.csv", "adf.csv", "kpss.csv",
)


def _diagnostics(root, identity):
    files = [root / "artifacts/diagnostics" / name for name in DIAGNOSTICS]
    files += [root / identity["run_dir"] / "selection" / name for name in ("selection_ledger.csv", "selection_summary.json")]
    missing = [str(p.relative_to(root)) for p in files if not p.is_file() or p.stat().st_size == 0]
    if missing:
        raise MissingEvidence("Missing diagnostic/selection evidence: " + ", ".join(missing))
    return {"files_present": len(files), "interpretation_reviewed": False}


def _dashboard(root):
    path = root / "artifacts/tests/dashboard_browser.json"
    if not path.is_file():
        raise MissingEvidence("Browser user-path evidence is missing")
    receipt = _json(path)
    if receipt.get("passed") is not True or receipt.get("views") != 5 or receipt.get("exceptions"):
        raise ValueError("Dashboard browser checks failed")
    hashes = receipt.get("file_hashes", {})
    required = {"app/streamlit_app.py", "scripts/verify_dashboard.cjs",
                "artifacts/tests/dashboard-desktop.png", "artifacts/tests/dashboard-mobile.png"}
    if not required.issubset(hashes):
        raise MissingEvidence("Browser receipt does not bind current UI and screenshots")
    for relative, expected in hashes.items():
        target = _safe(root, relative)
        if not target.is_file() or _file_hash(target) != expected:
            raise ValueError("Dashboard browser evidence is stale")
    return {"views": 5, "desktop": receipt["desktop"], "mobile": receipt["mobile"]}


def review_basis(root, config):
    """Return a current evidence basis without claiming inspection or approval."""
    root = Path(root).resolve()
    identity = _identity(root, config)
    directories = ("src/gridshift", "specifications", "artifacts/diagnostics",
                   identity["run_dir"], "artifacts/sensitivities", "artifacts/deployment",
                   "reports", "slides", "notebooks")
    files = {p for directory in directories for p in (root / directory).rglob("*")
             if p.is_file() and "__pycache__" not in p.parts and p.suffix != ".pyc"}
    files.update(root / name for name in (
        "app/streamlit_app.py", "scripts/verify_dashboard.cjs", "pyproject.toml", "uv.lock",
        "PROTOCOL_AMENDMENTS.md", "data/manifest.json", "artifacts/results.json",
        "artifacts/protocol_freeze.json", "artifacts/tests/core.xml",
        "artifacts/tests/core_receipt.json", "artifacts/tests/dashboard_browser.json",
    ))
    files.update((root / "artifacts/tests").glob("dashboard-*.png"))
    missing = [str(p) for p in files if not p.is_file()]
    if missing:
        raise MissingEvidence("Review basis lacks required evidence: " + ", ".join(missing))
    payload = {
        "identity": {key: identity[key] for key in ("study_id", "data_hash", "protocol_hash", "software_hash")},
        "file_hashes": {p.relative_to(root).as_posix(): _file_hash(p) for p in sorted(files)},
    }
    return {**payload, "basis_sha256": _hash(_json_bytes(payload))}


def record_review(root, config, *, role, reviewer_id, basis_sha256, decisions, notes):
    """Persist an explicit reviewer attestation; never infer approvals."""
    if role not in {"root", "independent"} or not isinstance(reviewer_id, str) or not reviewer_id.strip():
        raise ValueError("A declared root or independent reviewer identity is required")
    if set(decisions) != REVIEW_REQUIRED or any(v not in {"PASS", "FAIL", "PARTIAL"} for v in decisions.values()):
        raise ValueError("Every review-required requirement needs an explicit verdict")
    if not isinstance(notes, str) or not notes.strip():
        raise ValueError("Review notes must state evidence, findings, and limitations")
    basis = review_basis(root, config)
    if basis["basis_sha256"] != basis_sha256:
        raise ValueError("Review basis changed since inspection")
    receipt = {**basis, "role": role, "reviewer_id": reviewer_id,
               "reviewed_at_utc": datetime.now(timezone.utc).isoformat(),
               "decisions": decisions, "notes": notes,
               "scope": "Declared review attestation; reviewer identity is not cryptographically authenticated"}
    _write(Path(root) / "artifacts/reviews" / f"{role}.json", json.dumps(receipt, indent=2, sort_keys=True) + "\n")
    return receipt


def _review_evidence(root, config):
    try:
        basis = review_basis(root, config)
    except Exception as exc:
        return {"status": "PARTIAL", "details": str(exc), "decisions": {}}
    receipts = []
    for role in ("root", "independent"):
        path = root / "artifacts/reviews" / f"{role}.json"
        if not path.is_file():
            return {"status": "PARTIAL", "details": f"Missing {role} review", "decisions": {}, "basis_sha256": basis["basis_sha256"]}
        try:
            receipt = _json(path)
            if (receipt.get("role") != role or not receipt.get("reviewer_id")
                or receipt.get("basis_sha256") != basis["basis_sha256"]
                or receipt.get("file_hashes") != basis["file_hashes"]
                or receipt.get("identity") != basis["identity"]
                or set(receipt.get("decisions", {})) != REVIEW_REQUIRED
                or any(v not in {"PASS", "FAIL", "PARTIAL"} for v in receipt["decisions"].values())):
                raise ValueError("Stale or invalid review")
            receipts.append(receipt)
        except Exception as exc:
            return {"status": "PARTIAL", "details": f"{role}: {exc}", "decisions": {}}
    if receipts[0]["reviewer_id"] == receipts[1]["reviewer_id"]:
        return {"status": "PARTIAL", "details": "Reviewers must have distinct declared identities", "decisions": {}}
    decisions = {req: _req_status([{"status": r["decisions"][req]} for r in receipts]) for req in REVIEW_REQUIRED}
    return {"status": _req_status([{"status": value} for value in decisions.values()]),
            "decisions": decisions, "details": "Both current declared reviews match the evidence basis",
            "basis_sha256": basis["basis_sha256"], "reviewers": [r["reviewer_id"] for r in receipts]}


def _check(name, action, checks):
    try:
        checks[name] = {"status": "PASS", "details": action()}
    except MissingEvidence as exc:
        checks[name] = {"status": "PARTIAL", "details": str(exc)}
    except Exception as exc:
        checks[name] = {"status": "FAIL", "details": f"{type(exc).__name__}: {exc}"}
    return checks[name]


def _req_status(records):
    states = {r.get("status", "PARTIAL") for r in records}
    return "FAIL" if "FAIL" in states else "PARTIAL" if "PARTIAL" in states else "PASS"


def _write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".tmp")
    temp.write_text(text, encoding="utf-8", newline="\n")
    os.replace(temp, path)


def _requirements(checks, identity, review_evidence=None):
    rows = []
    for req_id, title, links in REQS:
        records = [checks.get(key, {"status": "PARTIAL", "details": "Not checked"}) for key in links]
        status = _req_status(records) if records else "PARTIAL"
        evidence = "; ".join(f"{key}: {r.get('details')}" for key, r in zip(links, records))
        next_action = "None for this executable check."
        if req_id in REVIEW_REQUIRED:
            review = review_evidence or {}
            review_status = review.get("decisions", {}).get(req_id, "PARTIAL")
            status = _req_status([{"status": status}, {"status": review_status}])
            evidence += "; review: " + review.get("details", "Fresh scientific/content/visual review required")
            if review_status != "PASS":
                next_action = "Complete current root and independent reviews against review_basis()."
        if status != "PASS" and next_action == "None for this executable check.":
            next_action = "Produce or repair the named evidence, then rerun offline verification."
        rows.append({"id": req_id, "title": title, "status": status, "evidence": evidence, "next_action": next_action})
    if identity is None:
        for row in rows:
            if row["status"] == "PASS":
                row["status"], row["next_action"] = "PARTIAL", "Complete compatible freeze/data/results identity first."
    return rows


def verify(root, config):
    """Check frozen local evidence without scoring or modifying experiment outputs."""
    root, config = Path(root).resolve(), dict(config)
    checks = {}
    identity_check = _check("identity", lambda: _identity(root, config), checks)
    identity = identity_check["details"] if identity_check["status"] == "PASS" else None
    _check("data", lambda: _data(root, identity or {}), checks)
    _check("dashboard", lambda: _dashboard(root), checks)
    for name, action in (
        ("forecasts", lambda: _forecasts(root, config, identity)),
        ("primary", lambda: _primary(root, config, identity)),
        ("queue", lambda: _queue(root, config, identity)),
        ("exports", lambda: _exports(root, identity)),
        ("notebook", lambda: _notebook(root)),
        ("deployment", lambda: _deployment(root, identity)),
        ("diagnostics", lambda: _diagnostics(root, identity)),
    ):
        if identity is None:
            checks[name] = {"status": "PARTIAL", "details": "Blocked until compatible freeze/data/results identity exists"}
        else:
            _check(name, action, checks)
    if identity is None:
        checks["tests"] = {"status": "PARTIAL", "details": "Blocked until frozen software identity exists"}
    else:
        _check("tests", lambda: _tests(root, identity), checks)

    reviews = _review_evidence(root, config) if identity is not None else {}
    requirements = _requirements(checks, identity, reviews)
    states = [row["status"] for row in requirements]
    overall = "FAIL" if "FAIL" in states else "PARTIAL" if "PARTIAL" in states else "PASS"
    result = {
        "schema_version": 1, "overall_status": overall,
        "scope": "Bounded offline artifact checks; not scientific certification.",
        "fresh_review_required": reviews.get("status") != "PASS", "identity": identity, "checks": checks,
        "review_evidence": reviews,
        "requirements": requirements,
        "summary": {s.lower(): states.count(s) for s in ("PASS", "PARTIAL", "FAIL")},
    }
    _write(root / "artifacts/verification.json", json.dumps(result, indent=2, sort_keys=True, default=str, allow_nan=False) + "\n")
    lines = [
        "# GridShift Completion Checklist", "",
        f"Verification status: **{overall}**. Executable checks and declared review attestations are separate evidence. This does not certify actual avoided emissions or public deployment.",
        "", "| ID | Requirement | Status | Evidence | Next action |", "| --- | --- | --- | --- | ---",
    ]
    for row in requirements:
        cells = [row[k].replace("|", "\\|").replace("\n", " ") for k in ("id", "title", "status", "evidence", "next_action")]
        lines.append("| " + " | ".join(cells) + " |")
    lines += ["", "Machine-readable evidence: `artifacts/verification.json`.", ""]
    _write(root / "COMPLETION_CHECKLIST.md", "\n".join(lines))
    return result
